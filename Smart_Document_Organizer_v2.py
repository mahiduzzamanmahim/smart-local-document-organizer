import os
import shutil
import tkinter as tk
from tkinter import filedialog, messagebox
import re
import threading
import queue
import json
from datetime import datetime

# ============================================================
# SMART DOCUMENT ORGANIZER - PRODUCTION BUILD
# Local-only classification + threaded UI + undo history
# ============================================================

CATEGORY_KEYWORDS = {
    "Invoices": {
        "invoice": 10, "receipt": 8, "tax invoice": 12,
        "amount due": 10, "total due": 9, "payment due": 8,
        "vat": 5, "bill": 5
    },
    "Resumes": {
        "resume": 12, "curriculum vitae": 12, " cv ": 10,
        "professional summary": 8, "work experience": 6,
        "education": 4, "skills": 4, "employment": 5
    },
    "Research": {
        "research": 9, "abstract": 7, "methodology": 8,
        "literature review": 9, "references": 5, "doi": 10,
        "dataset": 6, "machine learning": 8, "deep learning": 8,
        "results": 4, "discussion": 4
    },
    "Study": {
        "assignment": 9, "homework": 9, "lecture": 7, "quiz": 8,
        "exam": 8, "course": 6, "semester": 6, "tutorial": 6,
        "problem set": 8, "notes": 6, "iitm": 10, "iit madras": 10
    },
    "Finance": {
        "bank": 6, "bank statement": 10, "transaction": 6,
        "account": 5, "balance": 6, "salary": 7, "financial": 7
    },
    "Legal": {
        "agreement": 8, "contract": 10, "legal": 9, "court": 10,
        "petition": 10, "affidavit": 10, "notice": 6,
        "terms and conditions": 7
    },
    "Documents": {
        "document": 3, "report": 4, "letter": 4,
        "application": 5, "form": 4
    }
}

IMAGE_EXTENSIONS = {
    ".jpg", ".jpeg", ".png", ".gif", ".bmp", ".webp", ".tiff", ".svg"
}

DOCUMENT_EXTENSIONS = {
    ".doc", ".docx", ".txt", ".rtf", ".md", ".odt"
}

SPREADSHEET_EXTENSIONS = {
    ".xls", ".xlsx", ".csv"
}

CATEGORY_FOLDERS = {
    "Invoices", "Resumes", "Research", "Study", "Finance",
    "Legal", "Documents", "Images", "Miscellaneous"
}

# History is kept beside the application and never leaves the computer.
LOG_FILE = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "organization_history.json"
)


# ============================================================
# CLASSIFICATION
# ============================================================

def normalize_text(text):
    text = text.lower()
    text = re.sub(r"[^a-z0-9\s]", " ", text)
    text = re.sub(r"\s+", " ", text)
    return " " + text.strip() + " "


NORMALIZED_CATEGORY_KEYWORDS = {
    category: [
        (normalize_text(keyword), points)
        for keyword, points in keywords.items()
    ]
    for category, keywords in CATEGORY_KEYWORDS.items()
}


def classify_from_filename(filename):
    extension = os.path.splitext(filename)[1].lower()
    name = normalize_text(filename)

    if extension in IMAGE_EXTENSIONS:
        return "Images", 100

    scores = {}
    for category, keywords in NORMALIZED_CATEGORY_KEYWORDS.items():
        score = 0
        for keyword_normalized, points in keywords:
            if keyword_normalized in name:
                score += points * 2
        scores[category] = score

    best_category = max(scores, key=scores.get)
    best_score = scores[best_category]

    if best_score >= 10:
        return best_category, best_score

    if (
        extension in DOCUMENT_EXTENSIONS
        or extension in SPREADSHEET_EXTENSIONS
        or extension == ".pdf"
    ):
        if best_score >= 5:
            return best_category, best_score

    return None, 0


def extract_text(file_path):
    extension = os.path.splitext(file_path)[1].lower()

    if extension in [".txt", ".md", ".rtf"]:
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read(30000)
        except Exception:
            return ""

    if extension == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(file_path)
            text = ""
            for page in reader.pages[:3]:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
            return text[:30000]
        except Exception:
            return ""

    if extension == ".docx":
        try:
            from docx import Document
            document = Document(file_path)
            text = ""
            for paragraph in document.paragraphs[:100]:
                text += paragraph.text + "\n"
            return text[:30000]
        except Exception:
            return ""

    if extension == ".xlsx":
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(
                file_path,
                read_only=True,
                data_only=True
            )
            text = ""
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows(max_row=30, max_col=20):
                    for cell in row:
                        if cell.value is not None:
                            text += str(cell.value) + " "
            return text[:30000]
        except Exception:
            return ""

    return ""


def classify_from_content(content):
    if not content:
        return "Miscellaneous", 0

    content = normalize_text(content)
    scores = {}

    for category, keywords in NORMALIZED_CATEGORY_KEYWORDS.items():
        score = 0
        for keyword_normalized, points in keywords:
            occurrences = content.count(keyword_normalized)
            if occurrences > 0:
                occurrences = min(occurrences, 5)
                score += points * occurrences
        scores[category] = score

    best_category = max(scores, key=scores.get)
    best_score = scores[best_category]

    if best_score >= 5:
        return best_category, best_score

    return "Miscellaneous", best_score


def classify_file(file_path):
    filename = os.path.basename(file_path)
    category, score = classify_from_filename(filename)

    if category is not None:
        return category, score, "Filename"

    content = extract_text(file_path)
    category, score = classify_from_content(content)
    return category, score, "Content"


# ============================================================
# SAFE FILE MOVEMENT
# ============================================================

def safe_move(file_path, destination_folder):
    """Move a file without overwriting an existing file.

    Returns the actual destination path so the move can be undone later.
    """
    filename = os.path.basename(file_path)
    destination = os.path.join(destination_folder, filename)

    if os.path.exists(destination):
        base, extension = os.path.splitext(filename)
        counter = 1

        while True:
            new_filename = f"{base}_{counter}{extension}"
            destination = os.path.join(destination_folder, new_filename)

            if not os.path.exists(destination):
                break

            counter += 1

    shutil.move(file_path, destination)
    return destination


# ============================================================
# ORGANIZATION HISTORY / UNDO
# ============================================================

def load_history():
    """Load local organization history safely."""
    if not os.path.exists(LOG_FILE):
        return []

    try:
        with open(LOG_FILE, "r", encoding="utf-8") as f:
            history = json.load(f)
            return history if isinstance(history, list) else []
    except (OSError, json.JSONDecodeError):
        return []


def save_history(history):
    """Atomically save organization history."""
    temp_file = LOG_FILE + ".tmp"

    with open(temp_file, "w", encoding="utf-8") as f:
        json.dump(history, f, indent=4, ensure_ascii=False)
        f.flush()
        os.fsync(f.fileno())

    os.replace(temp_file, LOG_FILE)


def start_history_operation(source_dir):
    """Create and persist a new operation before files are moved."""
    history = load_history()

    operation = {
        "timestamp": datetime.now().isoformat(timespec="seconds"),
        "source_directory": os.path.abspath(source_dir),
        "status": "in_progress",
        "moves": [],
        "created_directories": []
    }

    history.append(operation)
    save_history(history)
    return len(history) - 1


def record_move(operation_index, original_path, new_path):
    """Persist one successful move immediately."""
    history = load_history()

    if not (0 <= operation_index < len(history)):
        return

    history[operation_index]["moves"].append({
        "original": os.path.abspath(original_path),
        "new": os.path.abspath(new_path)
    })

    save_history(history)


def record_created_directory(operation_index, directory_path):
    """Remember directories created by this operation."""
    history = load_history()

    if not (0 <= operation_index < len(history)):
        return

    history[operation_index]["created_directories"].append(
        os.path.abspath(directory_path)
    )

    save_history(history)


def finish_history_operation(operation_index, status="completed"):
    """Mark the latest operation as completed or partially completed."""
    history = load_history()

    if not (0 <= operation_index < len(history)):
        return

    history[operation_index]["status"] = status
    history[operation_index]["finished_at"] = datetime.now().isoformat(
        timespec="seconds"
    )

    save_history(history)


def undo_last_organization():
    """Undo the most recent completed or partially completed organization."""
    history = load_history()

    if not history:
        messagebox.showinfo(
            "Undo",
            "There is no organization history to undo."
        )
        return

    operation = history[-1]
    moves = operation.get("moves", [])

    if not moves:
        history.pop()
        save_history(history)
        messagebox.showinfo(
            "Undo",
            "The last organization did not move any files."
        )
        return

    status = operation.get("status", "unknown")

    confirm = messagebox.askyesno(
        "Undo Last Organization",
        (
            f"Undo the last organization?\n\n"
            f"Date: {operation.get('timestamp', 'Unknown')}\n"
            f"Status: {status}\n"
            f"Files moved: {len(moves)}\n\n"
            "The files will be restored to their original locations."
        )
    )

    if not confirm:
        return

    restored = 0
    skipped = 0

    # Reverse order prevents path conflicts when several files were moved.
    for move in reversed(moves):
        original = move["original"]
        new = move["new"]

        # The file may have been manually moved/deleted after organization.
        if not os.path.exists(new):
            skipped += 1
            continue

        # Never overwrite a file that appeared at the original location.
        if os.path.exists(original):
            skipped += 1
            continue

        try:
            os.makedirs(os.path.dirname(original), exist_ok=True)
            shutil.move(new, original)
            restored += 1
        except Exception:
            skipped += 1

    # Remove directories created by the operation if they are now empty.
    removed_dirs = 0
    for directory in sorted(
        operation.get("created_directories", []),
        key=len,
        reverse=True
    ):
        try:
            if os.path.isdir(directory) and not os.listdir(directory):
                os.rmdir(directory)
                removed_dirs += 1
        except OSError:
            pass

    # If everything was restored, remove the history entry.
    if restored == len(moves) and skipped == 0:
        history.pop()
        save_history(history)
        result = (
            f"Undo complete.\n\n"
            f"Restored: {restored}\n"
            f"Empty folders removed: {removed_dirs}"
        )
    else:
        # Keep only the files that still need attention in the history.
        remaining_moves = []

        for move in moves:
            if os.path.exists(move["new"]):
                remaining_moves.append(move)

        operation["moves"] = remaining_moves
        operation["status"] = "partially_undone"
        history[-1] = operation
        save_history(history)

        result = (
            f"Undo partially completed.\n\n"
            f"Restored: {restored}\n"
            f"Skipped: {skipped}\n"
            f"Empty folders removed: {removed_dirs}\n\n"
            "Files that could not be restored remain in the history."
        )

    messagebox.showinfo("Undo Complete", result)


# ============================================================
# SAFETY LOGIC & THREADED PIPELINE
# ============================================================

WINDOWS_SYSTEM_DIRS = {
    "Windows",
    "Program Files",
    "Program Files (x86)",
    "ProgramData",
    "$Recycle.Bin",
    "Recovery",
    "System Volume Information",
    "Users",
}


def is_drive_root(path):
    abs_path = os.path.abspath(path)
    drive, _ = os.path.splitdrive(abs_path)
    root_path = os.path.join(drive, os.sep)
    return abs_path == root_path


def collect_files_to_process(root_dir):
    files_to_process = []

    for current_root, dirs, files in os.walk(root_dir):
        dirs[:] = [
            d for d in dirs
            if not d.startswith(".")
            and not d.startswith("$")
            and d not in CATEGORY_FOLDERS
            and d not in WINDOWS_SYSTEM_DIRS
        ]

        for filename in files:
            if filename.startswith(".") or filename.startswith("$"):
                continue

            if filename.lower() in [
                "desktop.ini",
                "ntuser.dat",
                "dumpstack.log",
                "thumbs.db",
                "organizer.py",
                "app.py",
                "organizer.exe",
                "newapp3.exe",
            ]:
                continue

            files_to_process.append(
                os.path.join(current_root, filename)
            )

    return files_to_process


def run_organization_thread(source_dir, ui_events):
    operation_index = None
    moves_completed = 0

    try:
        files_to_process = collect_files_to_process(source_dir)
        total = len(files_to_process)

        if total == 0:
            ui_events.put(("empty", None))
            return

        classified_files = []
        preview_counts = {}

        for index, file_path in enumerate(files_to_process, start=1):
            if index == 1 or index % 10 == 0 or index == total:
                ui_events.put(
                    (
                        "progress",
                        f"Analyzing {index}/{total}: "
                        f"{os.path.basename(file_path)}"
                    )
                )

            category, _, _ = classify_file(file_path)

            if category is None:
                category = "Miscellaneous"

            classified_files.append((file_path, category))
            preview_counts[category] = (
                preview_counts.get(category, 0) + 1
            )

        if total > 200 or is_drive_root(source_dir):
            preview_lines = [
                f"{category}: {count}"
                for category, count in sorted(preview_counts.items())
            ]

            preview_text = "\n".join(preview_lines[:8])

            if len(preview_lines) > 8:
                preview_text += "\n..."

            response = {"confirmed": False}
            confirmation_ready = threading.Event()

            ui_events.put(
                (
                    "confirm",
                    (
                        f"This folder contains {total} files.\n\n"
                        f"Preview:\n{preview_text}\n\n"
                        "Proceed with organization?",
                        response,
                        confirmation_ready,
                    ),
                )
            )

            confirmation_ready.wait()

            if not response["confirmed"]:
                ui_events.put(("cancelled", None))
                return

        # Start persistent history BEFORE moving any files.
        operation_index = start_history_operation(source_dir)

        category_count = {}
        created_directories = set()

        for index, (file_path, category) in enumerate(
            classified_files,
            start=1
        ):
            if index == 1 or index % 10 == 0 or index == total:
                ui_events.put(
                    (
                        "progress",
                        f"Moving {index}/{total}: "
                        f"{os.path.basename(file_path)}"
                    )
                )

            destination_folder = os.path.join(
                source_dir,
                category
            )

            if not os.path.exists(destination_folder):
                os.makedirs(destination_folder, exist_ok=True)
                created_directories.add(destination_folder)
                record_created_directory(
                    operation_index,
                    destination_folder
                )

            new_path = safe_move(
                file_path,
                destination_folder
            )

            # Persist each successful move immediately.
            record_move(
                operation_index,
                file_path,
                new_path
            )

            moves_completed += 1
            category_count[category] = (
                category_count.get(category, 0) + 1
            )

        finish_history_operation(
            operation_index,
            "completed"
        )

        summary = (
            f"Organized {moves_completed} total files:\n\n"
        )

        for cat, count in category_count.items():
            summary += f"📁 {cat}: {count}\n"

        summary += "\nYou can undo this operation from the main window."

        ui_events.put(("complete", summary))

    except Exception as e:
        # Preserve all successfully completed moves if something fails.
        if operation_index is not None:
            try:
                finish_history_operation(
                    operation_index,
                    "partial"
                )
            except Exception:
                pass

        ui_events.put(
            (
                "error",
                f"{str(e)}\n\n"
                f"Successfully moved before the error: {moves_completed}\n"
                "Use 'Undo Last Organization' to restore them."
            )
        )

    finally:
        ui_events.put(("enable_buttons", None))


# ============================================================
# UI EVENT PROCESSING
# ============================================================

def process_ui_events(
    root_window,
    status_label,
    progress_label,
    organize_button,
    undo_button
):
    while True:
        try:
            event_type, payload = ui_events.get_nowait()
        except queue.Empty:
            break

        if event_type == "progress":
            progress_label.config(text=payload)

        elif event_type == "empty":
            status_label.config(
                text="No files found to organize.",
                fg="black"
            )
            progress_label.config(text="")

        elif event_type == "confirm":
            prompt, response, confirmation_ready = payload

            response["confirmed"] = messagebox.askyesno(
                "Preview before organizing",
                prompt
            )

            confirmation_ready.set()

        elif event_type == "cancelled":
            status_label.config(
                text="Organization cancelled.",
                fg="black"
            )

        elif event_type == "complete":
            status_label.config(
                text="Success! Organization complete.",
                fg="green"
            )
            progress_label.config(text="Complete.")
            messagebox.showinfo("Success", payload)

        elif event_type == "error":
            status_label.config(
                text="Organization stopped because of an error.",
                fg="red"
            )
            messagebox.showerror(
                "Runtime Error",
                f"An error occurred:\n\n{payload}"
            )

        elif event_type == "enable_buttons":
            organize_button.config(state=tk.NORMAL)
            undo_button.config(state=tk.NORMAL)

    root_window.after(
        50,
        process_ui_events,
        root_window,
        status_label,
        progress_label,
        organize_button,
        undo_button,
    )


# ============================================================
# START ORGANIZATION
# ============================================================

def start_organization(
    source_dir,
    status_label,
    progress_label,
    organize_button,
    undo_button
):
    if not source_dir:
        messagebox.showerror(
            "Selection Error",
            "Please select a target folder first."
        )
        return

    if not os.path.isdir(source_dir):
        messagebox.showerror(
            "Selection Error",
            "The selected path is not a valid folder."
        )
        return

    if is_drive_root(source_dir):
        confirm = messagebox.askyesno(
            "Confirm Drive Root",
            "You selected a drive root such as C:\\ or D:\\. "
            "This may include a large number of files and system folders.\n\n"
            "Do you want to continue?"
        )

        if not confirm:
            status_label.config(
                text="Organization cancelled.",
                fg="black"
            )
            return

    organize_button.config(state=tk.DISABLED)
    undo_button.config(state=tk.DISABLED)

    status_label.config(
        text="Running background parser...",
        fg="blue"
    )
    progress_label.config(text="")

    threading.Thread(
        target=run_organization_thread,
        args=(source_dir, ui_events),
        daemon=True
    ).start()


# ============================================================
# BROWSE
# ============================================================

def browse_folder():
    folder_selected = filedialog.askdirectory()

    if folder_selected:
        folder_path_var.set(folder_selected)
        status_label.config(
            text="Target verified. Ready.",
            fg="black"
        )
        progress_label.config(text="")


# ============================================================
# USER INTERFACE
# ============================================================

root = tk.Tk()
root.title("Local Intelligent File Organizer")
root.geometry("600x380")
root.resizable(False, False)

# Optional application icon.
icon_path = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "Smart_Document_Organizer.ico"
)

if os.path.exists(icon_path):
    try:
        root.iconbitmap(icon_path)
    except tk.TclError:
        pass

folder_path_var = tk.StringVar()

header = tk.Label(
    root,
    text="Smart Document Organizer",
    font=("Arial", 17, "bold")
)
header.pack(pady=18)

frame = tk.Frame(root)
frame.pack(pady=10, fill="x")

entry = tk.Entry(
    frame,
    textvariable=folder_path_var,
    width=52
)
entry.pack(side="left", padx=10)

browse_btn = tk.Button(
    frame,
    text="Browse...",
    command=browse_folder
)
browse_btn.pack(side="left")

organize_btn = tk.Button(
    root,
    text="Organize Folder Now",
    bg="#0078D7",
    fg="white",
    font=("Arial", 11, "bold"),
    width=25,
    height=2,
    command=lambda: start_organization(
        folder_path_var.get(),
        status_label,
        progress_label,
        organize_btn,
        undo_btn
    )
)
organize_btn.pack(pady=14)

undo_btn = tk.Button(
    root,
    text="Undo Last Organization",
    width=25,
    command=undo_last_organization
)
undo_btn.pack(pady=4)

status_label = tk.Label(
    root,
    text="Select a directory to begin.",
    font=("Arial", 9, "italic")
)
status_label.pack(pady=8)

progress_label = tk.Label(
    root,
    text="",
    font=("Arial", 8)
)
progress_label.pack(pady=5)


if __name__ == "__main__":
    ui_events = queue.Queue()

    root.after(
        50,
        process_ui_events,
        root,
        status_label,
        progress_label,
        organize_btn,
        undo_btn,
    )

    root.mainloop()
